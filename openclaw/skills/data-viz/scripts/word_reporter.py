#!/usr/bin/env python3
"""
word_reporter.py — Word 文档自动生成器
支持：自动排版图表、生成目录、标题正文格式、页眉页脚

用法：
  python3 word_reporter.py --charts output/charts/ --output report.docx
  python3 word_reporter.py --charts charts/ --output report.docx --title "数据分析报告"
  python3 word_reporter.py --demo --output output/reports/demo_report.docx
"""

import argparse
import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


class WordReporter:
    def __init__(self, title="数据分析报告"):
        self.doc = Document()
        self.title = title
        self.charts = []
        self._setup_styles()

    def _setup_styles(self):
        """全局样式设置"""
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = "PingFang SC"
        font.size = Pt(11)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

        # 段落间距
        pf = style.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.35

        # 设置页边距
        for section in self.doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(3.17)
            section.right_margin = Cm(3.17)

    def add_title_page(self, subtitle=None, author=None, date=None):
        """封面页"""
        for _ in range(6):
            self.doc.add_paragraph()

        # 主标题
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(self.title)
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        run.font.name = "PingFang SC"

        # 副标题
        if subtitle:
            self.doc.add_paragraph()
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(subtitle)
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        # 信息
        info_items = []
        if author:
            info_items.append(f"作者：{author}")
        if date:
            info_items.append(f"日期：{date}")
        else:
            info_items.append(f"日期：{datetime.now().strftime('%Y年%m月%d日')}")

        self.doc.add_paragraph()
        for item in info_items:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(item)
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        self.doc.add_page_break()

    def add_heading(self, text, level=1):
        """添加标题"""
        h = self.doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.name = "PingFang SC"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        return h

    def add_paragraph(self, text, bold=False, indent=False, font_size=None):
        """添加正文段落"""
        p = self.doc.add_paragraph()
        if indent:
            p.paragraph_format.first_line_indent = Pt(22)
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "PingFang SC"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        if font_size:
            run.font.size = Pt(font_size)
        return p

    def add_image_with_caption(self, img_path, caption=None, width=5.5):
        """插入图片 + 图注"""
        if not os.path.exists(img_path):
            print(f"  ⚠️ 图片不存在: {img_path}")
            return

        # 居中的标题
        if caption:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"图：{caption}")
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.name = "PingFang SC"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

        # 图片
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width))

        # 图注说明
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"来源：本次数据采集与分析")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        run.font.name = "PingFang SC"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

        self.doc.add_paragraph()  # 空行

    def add_table(self, data, headers=None, caption=None):
        """插入表格"""
        if caption:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(f"表：{caption}")
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.name = "PingFang SC"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

        if headers:
            rows = [headers] + data
        else:
            rows = data

        if not rows:
            return

        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # 表头样式
        for j, cell in enumerate(table.rows[0].cells):
            cell.text = str(rows[0][j])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(10)
                    run.font.name = "PingFang SC"
                    run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
            # 表头背景色
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="2E86AB"/>')
            cell._tc.get_or_add_tcPr().append(shading)
            for r in cell.paragraphs[0].runs:
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # 数据行
        for i in range(1, len(rows)):
            for j, cell in enumerate(table.rows[i].cells):
                cell.text = str(rows[i][j])
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.font.name = "PingFang SC"
                        run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
                # 隔行变色
                if i % 2 == 0:
                    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F4FD"/>')
                    cell._tc.get_or_add_tcPr().append(shading)

        self.doc.add_paragraph()  # 空行

    def add_divider(self):
        """添加分隔线"""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("— — — — — — — — — —")
        run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        run.font.size = Pt(10)

    def add_bullet_list(self, items):
        """添加无序列表"""
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            run = p.add_run(item)
            run.font.name = "PingFang SC"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
            run.font.size = Pt(11)

    def add_numbered_list(self, items):
        """添加有序列表"""
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            run = p.add_run(f"{i}. {item}")
            run.font.name = "PingFang SC"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
            run.font.size = Pt(11)

    def add_charts_from_dir(self, charts_dir, intro_text=None):
        """从目录批量插入图表"""
        if not os.path.exists(charts_dir):
            print(f"⚠️ 图表目录不存在: {charts_dir}")
            return

        image_files = sorted([
            f for f in os.listdir(charts_dir)
            if f.lower().endswith((".png", ".jpg", ".jpeg", ".gif"))
        ])

        if not image_files:
            print("⚠️ 目录中无图片文件")
            return

        print(f"📷 发现 {len(image_files)} 张图表")

        if intro_text:
            self.add_paragraph(intro_text)

        for i, img_file in enumerate(image_files):
            img_path = os.path.join(charts_dir, img_file)
            # 从文件名提取标题
            caption = os.path.splitext(img_file)[0]
            # 清理标题
            caption = re.sub(r'^\d+[_-]', '', caption)
            caption = caption.replace('_', ' ').replace('-', ' ')
            self.add_image_with_caption(img_path, caption=caption)

    def add_summary_box(self, title, content):
        """添加总结/要点框（高亮段落）"""
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        run = p.add_run(f"📌 {title}：")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2E, 0x86, 0xAB)
        run.font.name = "PingFang SC"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

        run2 = p.add_run(content)
        run2.font.size = Pt(11)
        run2.font.name = "PingFang SC"
        run2.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

    def add_footer(self, text=None):
        """设置页脚"""
        section = self.doc.sections[0]
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if text:
            run = p.add_run(text)
        else:
            run = p.add_run("数据分析报告 — 自动生成")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    def add_page_number(self):
        """添加页码（页脚居中）"""
        section = self.doc.sections[0]
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 添加 "第 X 页"
        run = p.add_run("第 ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        # 页码域
        fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
        run2 = p.add_run()
        run2._r.append(fldChar1)

        instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
        run3 = p.add_run()
        run3._r.append(instrText)

        fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
        run4 = p.add_run()
        run4._r.append(fldChar2)

        run5 = p.add_run(" 页")
        run5.font.size = Pt(8)
        run5.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    def save(self, filepath):
        """保存文档"""
        dirpath = os.path.dirname(filepath)
        if dirpath:
            os.makedirs(dirpath, exist_ok=True)
        self.doc.save(filepath)
        print(f"✅ Word 文档已生成: {filepath}")


def generate_demo_report(output="output/reports/demo_report.docx"):
    """生成示例报告"""
    print("📝 生成示例报告...")
    reporter = WordReporter("数据分析综合报告")

    # 封面
    reporter.add_title_page(
        subtitle="基于 Python 的数据采集与可视化分析",
        author="数据可视化 Agent · AutoGen",
        date=datetime.now().strftime("%Y年%m月%d日")
    )

    # 目录
    reporter.add_heading("目录", level=1)
    toc_items = [
        "一、数据来源与采集方法",
        "二、数据清洗与预处理",
        "三、描述性统计分析",
        "四、可视化分析与洞察",
        "五、结论与建议",
    ]
    reporter.add_numbered_list(toc_items)
    reporter.doc.add_page_break()

    # 第一章
    reporter.add_heading("一、数据来源与采集方法", level=1)
    reporter.add_paragraph(
        "本次分析所使用的数据通过网络爬虫技术采集，主要来源为公开数据平台。"
        "数据采集使用 Python 的 requests 库和 BeautifulSoup 解析库，"
        "共采集到有效数据样本 1000+ 条。",
        indent=True
    )
    reporter.add_paragraph("主要数据字段包括：", bold=True)
    reporter.add_bullet_list([
        "月份：数据记录的时间维度",
        "类别：产品/服务品类分类",
        "销售额：各品类各月份的销售金额（万元）",
        "满意度评分：各城市用户满意度评价"
    ])

    # 第二章
    reporter.add_heading("二、数据清洗与预处理", level=1)
    reporter.add_paragraph(
        "数据清洗阶段主要进行了以下处理：缺失值检测与填充（均值法）、"
        "异常值识别与处理（3σ原则）、数据类型转换（日期标准化）、"
        "重复值去重等操作。最终获得干净可用的结构化数据。",
        indent=True
    )

    reporter.add_table(
        headers=["指标", "清洗前", "清洗后", "变化率"],
        data=[
            ["样本量", "1050", "1036", "-1.3%"],
            ["缺失值占比", "3.2%", "0%", "-100%"],
            ["异常值占比", "1.8%", "0.1%", "-94.4%"],
        ],
        caption="数据清洗前后对比"
    )

    # 第三章
    reporter.add_heading("三、描述性统计分析", level=1)
    reporter.add_paragraph(
        "对主要数值型指标进行描述性统计分析，结果如下表所示。"
        "从数据分布来看，各品类销售额差异明显，呈现一定的季节性波动特征。",
        indent=True
    )

    reporter.add_table(
        headers=["指标", "均值", "中位数", "标准差", "最小值", "最大值"],
        data=[
            ["销售额(万元)", "152.3", "148.5", "45.2", "78.0", "268.0"],
            ["满意度评分", "86.2", "87.0", "8.5", "70.0", "98.0"],
            ["市场份额(%)", "16.7", "16.0", "5.8", "8.0", "26.0"],
        ],
        caption="描述性统计结果"
    )

    # 第四章
    reporter.add_heading("四、可视化分析与洞察", level=1)
    reporter.add_paragraph(
        "以下为本次分析的核心可视化图表，每张图表都配有详细的解读说明。",
        indent=True
    )

    # 找示例图表
    demo_charts_dir = "output/charts"
    reporter.add_charts_from_dir(demo_charts_dir)

    # 图表解读
    reporter.add_heading("图表关键发现", level=2)
    reporter.add_summary_box("趋势发现",
        "从折线图可以看出，销售额呈现明显的季节性波动，"
        "下半年（尤其是Q4）为销售旺季，环比增长约25%。")
    reporter.add_summary_box("分布特征",
        "箱线图显示各品类销售额分布差异显著，"
        "电子产品类目的中位数和离散程度均最高。")
    reporter.add_summary_box("相关性",
        "热力图分析显示，各品类之间的销售额呈弱正相关关系，"
        "说明不同品类的增长驱动力相对独立。")

    # 第五章
    reporter.add_heading("五、结论与建议", level=1)
    reporter.add_paragraph(
        "基于本次数据采集与可视化分析，得出以下主要结论：",
        indent=True
    )
    reporter.add_numbered_list([
        "市场呈现明显的季节性特征，建议在Q3末提前布局Q4的营销资源。",
        "电子产品品类收入最高但波动最大，可考虑增加差异化产品线以分散风险。",
        "用户满意度普遍较高（均值86.2分），但部分城市仍存在提升空间。",
        "各品类增长驱动力独立，可针对不同品类制定个性化运营策略。"
    ])

    reporter.add_paragraph()
    reporter.add_paragraph(
        "未来可进一步拓展数据维度（如用户画像、竞品数据），"
        "并引入预测模型（时间序列、回归分析）以提升决策支持能力。",
        indent=True
    )

    # 页脚和页码
    reporter.add_footer()
    reporter.add_page_number()

    reporter.save(filepath=output)


def main():
    parser = argparse.ArgumentParser(description="📝 Word 报告生成器")
    parser.add_argument("--charts", help="图表目录（批量插入图片）")
    parser.add_argument("--output", default="output/reports/report.docx", help="输出 Word 文件路径")
    parser.add_argument("--title", default="数据分析报告", help="报告标题")
    parser.add_argument("--intro", default=None, help="图表章节引言")
    parser.add_argument("--demo", action="store_true", help="生成示例报告")
    args = parser.parse_args()

    if args.demo:
        generate_demo_report(args.output)
        return

    if not args.charts:
        print("请指定 --charts 或使用 --demo")
        parser.print_help()
        return

    reporter = WordReporter(args.title)
    reporter.add_title_page(date=datetime.now().strftime("%Y年%m月%d日"))

    reporter.add_heading("一、可视化分析", level=1)
    reporter.add_charts_from_dir(args.charts, intro_text=args.intro)

    reporter.add_footer()
    reporter.add_page_number()
    reporter.save(args.output)


if __name__ == "__main__":
    main()
